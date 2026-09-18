import io
import json
import os

import pandas as pd
from openpyxl.styles import Font, Alignment
from docx import Document


# ============================================================
# OUTPUT FOLDER
# ============================================================

OUTPUT_DIR = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "output"
)

os.makedirs(OUTPUT_DIR, exist_ok=True)


# ============================================================
# DATAFRAME
# ============================================================

def _df(records):
    return pd.DataFrame(records, columns=[
        "business_area",
        "requirement_group",
        "capability",
        "requirement",
        "page_number",
        "source_excerpt",
        "confidence"
    ])


# ============================================================
# EXCEL
# ============================================================

def build_excel(records):

    output_path = os.path.join(
        OUTPUT_DIR,
        "requirement_analysis.xlsx"
    )

    df = _df(records)

    summary = (
        df.groupby(
            "business_area",
            as_index=False
        )
        .size()
        .rename(
            columns={
                "size": "Requirement Count"
            }
        )
    )

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            sheet_name="Requirements",
            index=False
        )

        summary.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

        ws = writer.book["Requirements"]

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Header formatting
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                wrap_text=True
            )

        # Column widths
        widths = {
            "A": 25,
            "B": 30,
            "C": 30,
            "D": 60,
            "E": 15,
            "F": 60,
            "G": 15
        }

        for column, width in widths.items():
            ws.column_dimensions[column].width = width

        # Wrap text
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    wrap_text=True,
                    vertical="top"
                )

    print(
        f"Excel file saved: {output_path}"
    )

    return output_path


# ============================================================
# WORD
# ============================================================

def build_word(records):

    output_path = os.path.join(
        OUTPUT_DIR,
        "requirement_analysis.docx"
    )

    doc = Document()

    doc.add_heading(
        "AI Requirement Analysis Report",
        0
    )

    doc.add_paragraph(
        "Structured analysis generated from "
        "the submitted proposal/RFP."
    )

    current = None

    for r in records:

        if r["business_area"] != current:

            current = r["business_area"]

            doc.add_heading(
                current,
                level=1
            )

        doc.add_heading(
            r["requirement_group"],
            level=2
        )

        doc.add_paragraph(
            f"Capability: {r['capability']}"
        )

        doc.add_paragraph(
            r["requirement"]
        )

        doc.add_paragraph(
            f"Source: Page {r['page_number']} | "
            f"Confidence: {r['confidence']:.0%}"
        )

        # Add source excerpt
        if r.get("source_excerpt"):

            doc.add_paragraph(
                f"Source Excerpt: "
                f"{r['source_excerpt']}"
            )

    doc.save(output_path)

    print(
        f"Word file saved: {output_path}"
    )

    return output_path


# ============================================================
# JSON
# ============================================================

def build_json(records):

    output_path = os.path.join(
        OUTPUT_DIR,
        "requirement_analysis.json"
    )

    with open(
        output_path,
        "w",
        encoding="utf-8"
    ) as f:

        json.dump(
            records,
            f,
            indent=2,
            ensure_ascii=False
        )

    print(
        f"JSON file saved: {output_path}"
    )

    return output_path


# ============================================================
# BUILD ALL FILES
# ============================================================

def build_all_outputs(records):

    excel_path = build_excel(records)

    word_path = build_word(records)

    json_path = build_json(records)

    return {
        "excel": excel_path,
        "word": word_path,
        "json": json_path
    }


# ============================================================
# GET FILE BY TYPE
# ============================================================

def get_output_file(file_type):

    files = {
        "excel": os.path.join(
            OUTPUT_DIR,
            "requirement_analysis.xlsx"
        ),

        "word": os.path.join(
            OUTPUT_DIR,
            "requirement_analysis.docx"
        ),

        "json": os.path.join(
            OUTPUT_DIR,
            "requirement_analysis.json"
        )
    }

    return files.get(file_type)
